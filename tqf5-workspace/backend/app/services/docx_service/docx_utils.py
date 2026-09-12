# -*- coding: utf-8 -*-
"""
docx_utils.py
-------------
Helper module containing OpenXML utilities and formatting functions for python-docx.
Implements technical guidelines from tqf3_generation_guide.txt:
- Complex Script (CS) Thai Font setup (TH Sarabun New)
- Thai Native Dictionary Language tagging (th-TH) for word wrapping
- Borderless 2-column running footer with dynamic PAGE / NUMPAGES fields
- Table formatting (cantSplit, tblHeader, tblInd, fixed layout, shading)
- Paragraph alignment and thaiDistribute support
"""

import docx

from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START, WD_ORIENTATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    import pythainlp
    from pythainlp.corpus.common import thai_words
    from pythainlp.util import dict_trie
    from pythainlp.tokenize import Tokenizer
    
    # Custom dictionary to force breaking of long compound words in narrow tables
    custom_words = set(thai_words())
    custom_words.discard("องค์ประกอบพื้นฐาน")
    custom_words.add("องค์ประกอบ")
    custom_words.add("พื้นฐาน")
    custom_words.discard("สถาปัตยกรรม")
    custom_words.add("สถาปัตย")
    custom_words.add("กรรม")
    custom_words.discard("กระบวนการพัฒนา")
    custom_words.add("กระบวนการ")
    custom_words.add("พัฒนา")
    
    custom_trie = dict_trie(dict_source=custom_words)
    custom_tokenizer = Tokenizer(custom_dict=custom_trie, engine='newmm')
    HAS_PYTHAINLP = True
except ImportError:
    HAS_PYTHAINLP = False

def insert_zwsp_for_thai(text):
    if not text or not HAS_PYTHAINLP:
        return text
        
    words = custom_tokenizer.word_tokenize(text)
    
    result = ""
    for i, word in enumerate(words):
        if i > 0:
            prev_word = words[i-1]
            # Only insert ZWSP if neither the previous nor the current word is just whitespace
            # This prevents inserting ZWSP around spaces which breaks MS Word justification
            if prev_word.strip() != "" and word.strip() != "":
                result += '\u200b'
        result += word
    return result

def apply_zwsp_to_all_text(doc):
    if not HAS_PYTHAINLP:
        return

    def process_paragraph(p):
        for run in p.runs:
            if run.text:
                run.text = insert_zwsp_for_thai(run.text)

    for paragraph in doc.paragraphs:
        process_paragraph(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)

# Default styling constants
DEFAULT_FONT_NAME = "TH Sarabun New"
DEFAULT_CATEGORY_SIZE = 18
DEFAULT_HEADING_SIZE = 14
DEFAULT_BODY_SIZE = 14
DEFAULT_NOTE_SIZE = 12
HEADER_SHADING_HEX = "D9D9D9"


def setup_page_margins(doc):
    """
    Sets default page size to A4 Portrait with user specified margins:
    Top: 1.78 cm, Bottom: 1.02 cm, Left: 2.54 cm, Right: 2.54 cm.
    """
    for section in doc.sections:
        section.orientation = WD_ORIENTATION.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.78)
        section.bottom_margin = Cm(1.02)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)


def add_landscape_section(doc):
    """
    Adds a new A4 Landscape section (used for หมวดที่ 3) with margins:
    Top: 2.54 cm, Bottom: 2.54 cm, Left: 1.02 cm, Right: 1.78 cm.
    Printable width = 29.7 - 1.02 - 1.78 = 26.90 cm (~10.59 inches).
    """
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(1.02)
    section.right_margin = Cm(1.78)
    return section


def add_portrait_section(doc):
    """
    Adds a new A4 Portrait section returning to default margins:
    Top: 1.78 cm, Bottom: 1.02 cm, Left: 2.54 cm, Right: 2.54 cm.
    """
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.78)
    section.bottom_margin = Cm(1.02)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    return section



def format_run(run, font_name=DEFAULT_FONT_NAME, size_pt=14, bold=False, italic=False, color_rgb=None):
    """
    Formats a run with font specifications including OpenXML Complex Script (CS) tags.
    Ensures TH Sarabun New renders properly in Microsoft Word for Thai language text.
    - Sets rFonts for ascii, hAnsi, cs
    - Sets sz & szCs for font size
    - Sets b & bCs for Thai Complex Script bold weight
    - Sets i & iCs for Thai Complex Script italic style
    """
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()

    # 1. OpenXML rFonts for ascii, hAnsi, and cs (Complex Script for Thai)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

    # 2. Font size in half-points (sz & szCs)
    sz_val = str(int(size_pt * 2))
    for tag_name in ['sz', 'szCs']:
        node = rPr.find(qn(f'w:{tag_name}'))
        if node is None:
            node = OxmlElement(f'w:{tag_name}')
            rPr.append(node)
        node.set(qn('w:val'), sz_val)

    # 3. Bold tags (b & bCs) for Thai Complex Script font rendering
    b_tag = rPr.find(qn('w:b'))
    bCs_tag = rPr.find(qn('w:bCs'))
    if bold:
        if b_tag is None:
            b_tag = OxmlElement('w:b')
            rPr.append(b_tag)
        if bCs_tag is None:
            bCs_tag = OxmlElement('w:bCs')
            rPr.append(bCs_tag)
    else:
        if b_tag is not None:
            rPr.remove(b_tag)
        if bCs_tag is not None:
            rPr.remove(bCs_tag)

    # 4. Italic tags (i & iCs) for Thai Complex Script font rendering
    i_tag = rPr.find(qn('w:i'))
    iCs_tag = rPr.find(qn('w:iCs'))
    if italic:
        if i_tag is None:
            i_tag = OxmlElement('w:i')
            rPr.append(i_tag)
        if iCs_tag is None:
            iCs_tag = OxmlElement('w:iCs')
            rPr.append(iCs_tag)
    else:
        if i_tag is not None:
            rPr.remove(i_tag)
        if iCs_tag is not None:
            rPr.remove(iCs_tag)


def format_paragraph_runs(p, font_name=DEFAULT_FONT_NAME, size_pt=DEFAULT_BODY_SIZE, bold=False, italic=False, color_rgb=None):
    """
    Safely applies font formatting to all runs inside a paragraph.
    If paragraph has no runs, does nothing.
    """
    for r in p.runs:
        format_run(r, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)



def set_run_language_thai(run):
    """
    Applies OpenXML w:lang tag (th-TH) to a run.
    Enables Word's native Thai dictionary word-wrapping engine.
    """
    rPr = run._r.get_or_add_rPr()
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')


def apply_thai_language_to_all_runs(doc):
    """
    Applies th-TH language tag to all runs in document paragraphs and tables,
    AND sets document-level defaults so Word uses Thai dictionary word-wrapping globally.
    
    Two-pronged approach:
    1. Document-level: Set w:docDefaults rPrDefault lang w:val, w:bidi, and w:eastAsia to 'th-TH' and CS font to TH Sarabun New.
       This ensures Word applies Thai word-wrapping rules by default for all Complex Script text.
    2. Run-level: Tag every run with w:lang bidi='th-TH' as a fallback guarantee.
    """
    # --- 1. Document-level defaults ---
    styles_element = doc.styles.element
    doc_defaults = styles_element.find(qn('w:docDefaults'))
    if doc_defaults is not None:
        rPr_default = doc_defaults.find(qn('w:rPrDefault'))
        if rPr_default is not None:
            rPr = rPr_default.find(qn('w:rPr'))
            if rPr is not None:
                # Fix default language: set val, bidi, eastAsia to th-TH
                lang = rPr.find(qn('w:lang'))
                if lang is not None:
                    lang.set(qn('w:val'), 'th-TH')
                    lang.set(qn('w:bidi'), 'th-TH')
                    lang.set(qn('w:eastAsia'), 'th-TH')
                else:
                    lang = OxmlElement('w:lang')
                    lang.set(qn('w:val'), 'th-TH')
                    lang.set(qn('w:bidi'), 'th-TH')
                    lang.set(qn('w:eastAsia'), 'th-TH')
                    rPr.append(lang)

                # Fix default CS font: set to TH Sarabun New
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is not None:
                    rFonts.set(qn('w:cs'), DEFAULT_FONT_NAME)
                else:
                    rFonts = OxmlElement('w:rFonts')
                    rFonts.set(qn('w:cs'), DEFAULT_FONT_NAME)
                    rPr.append(rFonts)

                # Fix default CS font size (14pt = 28 half-points)
                szCs = rPr.find(qn('w:szCs'))
                if szCs is not None:
                    szCs.set(qn('w:val'), str(int(DEFAULT_BODY_SIZE * 2)))
                else:
                    szCs = OxmlElement('w:szCs')
                    szCs.set(qn('w:val'), str(int(DEFAULT_BODY_SIZE * 2)))
                    rPr.append(szCs)

    # --- 2. Run-level: tag every run with th-TH ---
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_run_language_thai(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_language_thai(run)


def set_thai_distributed(paragraph):
    """
    Sets paragraph alignment to thaiDistribute for justified Thai text.
    """
    pPr = paragraph._p.get_or_add_pPr()
    jc = pPr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        pPr.append(jc)
    jc.set(qn('w:val'), 'thaiDistribute')


def add_bottom_border_to_paragraph(p, color_hex="000000", sz="12"):
    """
    Adds a solid horizontal bottom border to a paragraph (used for header line separator).
    """
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)


def add_xml_field(run, field_name):

    """
    Inserts a dynamic OpenXML field code (e.g. PAGE, NUMPAGES) inside a run.
    """
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field_name

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def setup_header_footer(doc, course_code, course_name_th):
    """
    Configures a running footer using a borderless 2-column table:
    - Left column (~3.75 in): เค้าโครงรายวิชา [course_code] [course_name_th]
    - Right column (~2.52 in): หน้า PAGE/NUMPAGES
    Font size: 14pt TH Sarabun New, black color.
    Clears top header.
    """
    section = doc.sections[0]
    section.different_first_page_header_footer = False

    # 1. Clear top header
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.text = ""
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    # 2. Configure footer with borderless table
    footer = section.footer
    footer.is_linked_to_previous = False

    for p in footer.paragraphs:
        p.text = ""

    # Create table element
    tbl = OxmlElement('w:tbl')
    
    # Table properties
    tblPr = OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')  # 100% width
    tblPr.append(tblW)

    # Remove all borders
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        border.set(qn('w:sz'), '0')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    tblPr.append(tblBorders)

    # Zero margins
    tblCellMar = OxmlElement('w:tblCellMar')
    for side in ['top', 'left', 'bottom', 'right']:
        mar = OxmlElement(f'w:{side}')
        mar.set(qn('w:w'), '0')
        mar.set(qn('w:type'), 'dxa')
        tblCellMar.append(mar)
    tblPr.append(tblCellMar)
    tbl.append(tblPr)

    # Grid columns
    tblGrid = OxmlElement('w:tblGrid')
    col1 = OxmlElement('w:gridCol')
    col1.set(qn('w:w'), '5400')  # ~3.75 in
    tblGrid.append(col1)
    col2 = OxmlElement('w:gridCol')
    col2.set(qn('w:w'), '3629')  # ~2.52 in
    tblGrid.append(col2)
    tbl.append(tblGrid)

    # Table Row
    tr = OxmlElement('w:tr')

    # Left cell (Course title)
    tc_left = OxmlElement('w:tc')
    tcPr_l = OxmlElement('w:tcPr')
    tcW_l = OxmlElement('w:tcW')
    tcW_l.set(qn('w:w'), '5400')
    tcW_l.set(qn('w:type'), 'dxa')
    tcPr_l.append(tcW_l)
    tc_left.append(tcPr_l)

    p_left = OxmlElement('w:p')
    pPr_l = OxmlElement('w:pPr')
    jc_l = OxmlElement('w:jc')
    jc_l.set(qn('w:val'), 'left')
    pPr_l.append(jc_l)
    p_left.append(pPr_l)

    r_left = OxmlElement('w:r')
    t_left = OxmlElement('w:t')
    t_left.text = f"เค้าโครงรายวิชา {course_code} {course_name_th}"
    r_left.append(t_left)
    p_left.append(r_left)
    tc_left.append(p_left)
    tr.append(tc_left)

    # Right cell (Page numbers)
    tc_right = OxmlElement('w:tc')
    tcPr_r = OxmlElement('w:tcPr')
    tcW_r = OxmlElement('w:tcW')
    tcW_r.set(qn('w:w'), '3629')
    tcW_r.set(qn('w:type'), 'dxa')
    tcPr_r.append(tcW_r)
    tc_right.append(tcPr_r)

    p_right = OxmlElement('w:p')
    pPr_r = OxmlElement('w:pPr')
    jc_r = OxmlElement('w:jc')
    jc_r.set(qn('w:val'), 'right')
    pPr_r.append(jc_r)
    p_right.append(pPr_r)

    # "หน้า "
    r1 = OxmlElement('w:r')
    t1 = OxmlElement('w:t')
    t1.text = "หน้า "
    r1.append(t1)
    p_right.append(r1)

    # PAGE field
    r_page = OxmlElement('w:r')
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr1 = OxmlElement('w:instrText')
    instr1.set(qn('xml:space'), 'preserve')
    instr1.text = "PAGE"
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    r_page.append(fld1)
    r_page.append(instr1)
    r_page.append(fld2)
    r_page.append(fld3)
    p_right.append(r_page)

    # "/"
    r2 = OxmlElement('w:r')
    t2 = OxmlElement('w:t')
    t2.text = "/"
    r2.append(t2)
    p_right.append(r2)

    # NUMPAGES field
    r_numpages = OxmlElement('w:r')
    fld4 = OxmlElement('w:fldChar')
    fld4.set(qn('w:fldCharType'), 'begin')
    instr2 = OxmlElement('w:instrText')
    instr2.set(qn('xml:space'), 'preserve')
    instr2.text = "NUMPAGES"
    fld5 = OxmlElement('w:fldChar')
    fld5.set(qn('w:fldCharType'), 'separate')
    fld6 = OxmlElement('w:fldChar')
    fld6.set(qn('w:fldCharType'), 'end')
    r_numpages.append(fld4)
    r_numpages.append(instr2)
    r_numpages.append(fld5)
    r_numpages.append(fld6)
    p_right.append(r_numpages)

    tc_right.append(p_right)
    tr.append(tc_right)
    tbl.append(tr)

    # Add tbl to footer paragraph
    footer.paragraphs[0]._p.addprevious(tbl)

    # Format runs in footer
    for p in footer.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    # Apply font formatting to footer elements
    for element in tbl.iter():
        if element.tag == qn('w:r'):
            rPr = element.get_or_add_rPr()
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), DEFAULT_FONT_NAME)
            rFonts.set(qn('w:hAnsi'), DEFAULT_FONT_NAME)
            rFonts.set(qn('w:cs'), DEFAULT_FONT_NAME)
            rPr.append(rFonts)
            
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '28')  # 14pt
            szCs = OxmlElement('w:szCs')
            szCs.set(qn('w:val'), '28')
            rPr.append(sz)
            rPr.append(szCs)

            color = OxmlElement('w:color')
            color.set(qn('w:val'), '000000')
            rPr.append(color)

            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), 'th-TH')
            lang.set(qn('w:bidi'), 'th-TH')
            rPr.append(lang)


# -------------------------------------------------------------------------
# Table Helper Functions
# -------------------------------------------------------------------------

def set_cell_background(cell, hex_color):
    """
    Sets background shading color of a table cell.
    """
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_row_cant_split(row):
    """
    Prevents a table row from splitting across pages.
    """
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))


def set_row_tbl_header(row):
    """
    Repeats a table header row across page breaks.
    """
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))


def set_table_left_indent(table, indent_inches=0.0):
    """
    Sets table left indent using tblInd (in dxa).
    0.0 inch = 0 dxa (aligns table left edge flush with left margin).
    Removes w:jc center alignment if present so left indent is respected by Word.
    """
    tblPr = table._tbl.tblPr

    existing_jc = tblPr.find(qn('w:jc'))
    if existing_jc is not None:
        tblPr.remove(existing_jc)

    existing_ind = tblPr.find(qn('w:tblInd'))
    if existing_ind is not None:
        tblPr.remove(existing_ind)

    tblInd = OxmlElement('w:tblInd')
    dxa_val = str(int(indent_inches * 1440))
    tblInd.set(qn('w:w'), dxa_val)
    tblInd.set(qn('w:type'), 'dxa')
    tblPr.append(tblInd)




def set_table_borders(table, color="000000", sz="4", val="single"):
    """
    Sets clean explicit outer and inner grid borders on a table.
    color: Hex color (default '000000' black)
    sz: Border line width (4 = 0.5pt, 8 = 1pt)
    val: Border line style (default 'single')
    """
    tblPr = table._tbl.tblPr
    existing_bdr = tblPr.find(qn('w:tblBorders'))
    if existing_bdr is not None:
        tblPr.remove(existing_bdr)

    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)



def set_table_compact_cell_margins(table, top_pt=1.0, bottom_pt=1.0, left_pt=4.0, right_pt=4.0):
    """
    Sets compact cell margins (padding) inside table cells.
    top_pt, bottom_pt: Top and bottom margin in points (default 1.0 pt = 20 dxa).
    left_pt, right_pt: Left and right margin in points (default 4.0 pt = 80 dxa).
    """
    tblPr = table._tbl.tblPr
    existing_mar = tblPr.find(qn('w:tblCellMar'))
    if existing_mar is not None:
        tblPr.remove(existing_mar)

    tblCellMar = OxmlElement('w:tblCellMar')
    for side, pt_val in [('top', top_pt), ('bottom', bottom_pt), ('left', left_pt), ('right', right_pt)]:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(int(pt_val * 20)))
        node.set(qn('w:type'), 'dxa')
        tblCellMar.append(node)
    tblPr.append(tblCellMar)


def format_table_header(row, bg_hex=HEADER_SHADING_HEX):
    """
    Applies standard styling to a table header row:
    - Background color #D9D9D9
    - Bold 14pt text centered
    - cantSplit & tblHeader
    - Compact paragraph spacing (0.5pt before/after, 1.0 line spacing)
    """
    set_row_cant_split(row)
    set_row_tbl_header(row)

    for cell in row.cells:
        set_cell_background(cell, bg_hex)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                format_run(r, size_pt=DEFAULT_BODY_SIZE, bold=True)


def set_table_fixed_width(table, col_widths_inches=None):
    """
    Disables autofit and sets fixed column widths on a table.
    """
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    if col_widths_inches:
        for row in table.rows:
            for idx, width in enumerate(col_widths_inches):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)


def add_paragraph_styled(doc, text="", font_size=DEFAULT_BODY_SIZE, bold=False, italic=False,
                         space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT,
                         first_line_indent=0.0, left_indent=0.0, thai_distribute=False):
    """
    Utility function to create a paragraph with customized font, spacing, and alignment.
    """
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if first_line_indent > 0:
        p.paragraph_format.first_line_indent = Inches(first_line_indent)
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)

    if thai_distribute:
        set_thai_distributed(p)

    if text:
        r = p.add_run(text)
        format_run(r, size_pt=font_size, bold=bold, italic=italic)

    return p


def format_and_scale_all_data_tables(doc, indent_inches=0.25):
    """
    Applies left indent (0.25 in), compact cell padding, tight line spacing, explicit grid borders,
    and scales column widths according to section orientation (Portrait vs Landscape).
    Aligns table left border with subheading numbers (1., 2., 3.).
    - Portrait target width = 6.02 inches (6.27 - 0.25)
    - Landscape target width = 10.34 inches (10.59 - 0.25)
    """
    body_children = list(doc._body._element)
    tbl_index = 0

    for idx, elem in enumerate(body_children):
        if elem.tag.endswith('tbl'):
            if tbl_index == 0:
                # Table 0 is Page 1 Top Header table - skip scaling
                tbl_index += 1
                continue

            table = doc.tables[tbl_index]
            tbl_index += 1

            set_table_left_indent(table, indent_inches)
            set_table_borders(table, color="000000", sz="4", val="single")
            set_table_compact_cell_margins(table, top_pt=1.0, bottom_pt=1.0, left_pt=4.0, right_pt=4.0)

            # Reduce line spacing inside all cell paragraphs for compact table appearance
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(0.5)
                        p.paragraph_format.space_after = Pt(0.5)
                        p.paragraph_format.line_spacing = 1.0

            # Look ahead from current element position to find next section properties (sectPr)
            is_landscape = False
            for next_elem in body_children[idx:]:
                sectPr = None
                pPr = next_elem.find(qn('w:pPr'))
                if pPr is not None:
                    sectPr = pPr.find(qn('w:sectPr'))
                if sectPr is None and next_elem.tag.endswith('sectPr'):
                    sectPr = next_elem

                if sectPr is not None:
                    pgSz = sectPr.find(qn('w:pgSz'))
                    if pgSz is not None and pgSz.attrib.get(qn('w:orient')) == 'landscape':
                        is_landscape = True
                    break

            target_width_inches = 10.34 if is_landscape else 6.02

            # Find a row without horizontal merges (where len(row.cells) == len(table.columns))
            unmerged_row = None
            for r in table.rows:
                if len(r.cells) == len(table.columns):
                    unmerged_row = r
                    break

            if len(table.columns) == 16:
                # Table 4 Curriculum Mapping Matrix explicit widths (col 0 = 1.22 in, cols 1..15 = 0.32 in each)
                set_table_fixed_width(table, [1.22] + [0.32] * 15)
            elif len(table.columns) == 6 and table.rows[0].cells[0].paragraphs[0].text.strip() == "ลำดับที่":
                # Table 3 CLOs Table explicit narrow width (K, S, E, C expanded to 0.80 in so Knowledge stays on line 1)
                set_table_fixed_width(table, [0.65, 2.00, 0.80, 0.80, 0.80, 0.80])
            elif len(table.columns) == 7 and "ผลลัพธ์การเรียนรู้รายวิชา CLO" in table.rows[0].cells[0].paragraphs[0].text:
                # Item 8 Self Assessment Table explicit fixed widths (total = 6.02 in)
                set_table_fixed_width(table, [2.50, 0.40, 0.40, 0.40, 0.40, 0.40, 1.52])
            elif len(table.columns) == 7 and "งานประเมิน" in table.rows[0].cells[0].paragraphs[0].text:
                # Table 8 Evaluation Plan Table explicit fixed widths (total = 6.02 in)
                set_table_fixed_width(table, [1.72, 0.67, 0.77, 0.67, 0.67, 0.77, 0.75])
            elif unmerged_row is not None:
                widths = [c.width.inches for c in unmerged_row.cells if c.width and c.width.inches > 0]
                total_w = sum(widths)

                if total_w > 0 and abs(total_w - target_width_inches) > 0.05:
                    ratio = target_width_inches / total_w
                    new_widths = [w * ratio for w in widths]
                    set_table_fixed_width(table, new_widths)







